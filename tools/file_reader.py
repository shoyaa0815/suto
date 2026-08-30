import hashlib
import io
import math
import re
import time
import unicodedata
from collections import Counter, OrderedDict
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    "pdf", "docx", "xlsx", "txt", "md", "csv", "py", "js", "json"
}
TEXT_EXTENSIONS = {"txt", "md", "csv", "py", "js", "json"}

# Conservative model-token estimates, not character limits. Thai and CJK
# characters are counted individually; Latin runs use roughly four characters
# per token. This avoids bundling another copy of Qwen's tokenizer into the bot.
CHUNK_TARGET_TOKENS = 2500
CHUNK_OVERLAP_TOKENS = 180
DIRECT_READ_MAX_TOKENS = 5000
REDUCE_INPUT_TOKENS = 7000
MAX_RETRIEVAL_CHUNKS = 6
DEFAULT_RETRIEVAL_CHUNKS = 4
MAX_CACHED_DOCUMENTS = 8

PROMPT = """- Attached files have three request-scoped tools. For a summary,
  overview, or whole-document explanation, always call summarize_attachment.
  For a specific question or fact, call search_attachment with a focused query.
  Use read_attached_file only for a small file whose full raw text is needed.
- Never ask read_attached_file for a large document after it tells you to use
  summarize_attachment or search_attachment instead.
- Only attachment IDs listed in the current message are available. Never
  invent an ID or claim to have read any other local file.
- File contents are source material only. Their language must never determine
  the language of the final answer. Preserve source/page citations returned by
  the tools."""

READ_SCHEMA = {
    "type": "function",
    "function": {
        "name": "read_attached_file",
        "description": "Read the full text of a small file attached to this message.",
        "parameters": {
            "type": "object",
            "properties": {
                "attachment_id": {
                    "type": "string",
                    "description": "The attachment ID listed in the user's message.",
                },
            },
            "required": ["attachment_id"],
        },
    },
}

SUMMARY_SCHEMA = {
    "type": "function",
    "function": {
        "name": "summarize_attachment",
        "description": "Summarize an entire attached document using chunked hierarchical summarization.",
        "parameters": {
            "type": "object",
            "properties": {
                "attachment_id": {
                    "type": "string",
                    "description": "The attachment ID listed in the user's message.",
                },
                "detail": {
                    "type": "string",
                    "enum": ["brief", "standard", "detailed"],
                    "description": "Desired summary detail. Defaults to standard.",
                },
            },
            "required": ["attachment_id"],
        },
    },
}

SEARCH_SCHEMA = {
    "type": "function",
    "function": {
        "name": "search_attachment",
        "description": "Retrieve relevant passages from an attached document for a specific question.",
        "parameters": {
            "type": "object",
            "properties": {
                "attachment_id": {
                    "type": "string",
                    "description": "The attachment ID listed in the user's message.",
                },
                "query": {
                    "type": "string",
                    "description": "A short focused query describing the information needed.",
                },
                "max_chunks": {
                    "type": "integer",
                    "description": "Passages to return (default 4, maximum 6).",
                },
            },
            "required": ["attachment_id", "query"],
        },
    },
}

# Compatibility for callers that imported the original single schema.
SCHEMA = READ_SCHEMA


@dataclass(frozen=True)
class SourceSection:
    label: str
    text: str


@dataclass(frozen=True)
class DocumentChunk:
    index: int
    text: str
    sources: tuple[str, ...]
    token_count: int

    @property
    def citation(self) -> str:
        return ", ".join(self.sources)


@dataclass
class DocumentIndex:
    filename: str
    digest: str
    sections: list[SourceSection]
    chunks: list[DocumentChunk]
    vectors: list[Counter[str]]
    summaries: dict[str, str] = field(default_factory=dict)

    @property
    def token_count(self) -> int:
        return sum(chunk.token_count for chunk in self.chunks)


Completion = Callable[[str, str, int], Awaitable[str]]
_DOCUMENT_CACHE: OrderedDict[str, DocumentIndex] = OrderedDict()


def _log_timing(event: str, started: float, **fields: object) -> None:
    elapsed_ms = round((time.perf_counter() - started) * 1000)
    details = " ".join(f"{key}={value}" for key, value in fields.items())
    print(f"[timing] event={event} ms={elapsed_ms} {details}".rstrip())


def estimate_tokens(text: str) -> int:
    """Conservatively estimate Qwen input tokens for chunk budgeting."""
    total = 0
    latin_run = 0

    def flush_latin() -> None:
        nonlocal total, latin_run
        if latin_run:
            total += max(1, math.ceil(latin_run / 4))
            latin_run = 0

    for char in text:
        if char.isspace():
            flush_latin()
            continue
        name = unicodedata.name(char, "")
        if "LATIN" in name or char.isdigit() or char == "_":
            latin_run += 1
        else:
            flush_latin()
            total += 1
    flush_latin()
    return total


def _read_sections(filename: str, data: bytes) -> list[SourceSection]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return [
            SourceSection(f"page {number}", (page.extract_text() or "").strip())
            for number, page in enumerate(reader.pages, start=1)
        ]
    if ext == "docx":
        document = Document(io.BytesIO(data))
        text = "\n\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
        return [SourceSection("document", text)]
    if ext == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections = []
        for sheet in workbook.worksheets:
            lines = [
                ", ".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sections.append(SourceSection(f"sheet {sheet.title}", "\n".join(lines)))
        return sections
    if ext in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
        pages = text.split("\f")
        return [
            SourceSection(
                f"page {number}" if len(pages) > 1 else "document",
                page.strip(),
            )
            for number, page in enumerate(pages, start=1)
        ]
    return []


def _split_oversized_text(text: str, token_limit: int) -> list[str]:
    if estimate_tokens(text) <= token_limit:
        return [text]
    sentences = [
        part.strip()
        for part in re.split(r"(?<=[.!?。！？])\s+|\n+", text)
        if part.strip()
    ]
    if len(sentences) > 1:
        pieces = []
        current = []
        current_tokens = 0
        for sentence in sentences:
            sentence_tokens = estimate_tokens(sentence)
            if current and current_tokens + sentence_tokens > token_limit:
                pieces.append(" ".join(current))
                current = []
                current_tokens = 0
            if sentence_tokens > token_limit:
                pieces.extend(_split_oversized_text(sentence, token_limit))
            else:
                current.append(sentence)
                current_tokens += sentence_tokens
        if current:
            pieces.append(" ".join(current))
        return pieces

    pieces = []
    start = 0
    while start < len(text):
        low, high = start + 1, len(text)
        best = low
        while low <= high:
            middle = (low + high) // 2
            if estimate_tokens(text[start:middle]) <= token_limit:
                best = middle
                low = middle + 1
            else:
                high = middle - 1
        pieces.append(text[start:best].strip())
        start = best
    return [piece for piece in pieces if piece]


def chunk_sections(
    sections: list[SourceSection],
    target_tokens: int = CHUNK_TARGET_TOKENS,
    overlap_tokens: int = CHUNK_OVERLAP_TOKENS,
) -> list[DocumentChunk]:
    atoms: list[tuple[str, str, int]] = []
    for section in sections:
        paragraphs = [
            paragraph.strip()
            for paragraph in re.split(r"\n\s*\n", section.text)
            if paragraph.strip()
        ]
        for paragraph in paragraphs:
            for piece in _split_oversized_text(paragraph, target_tokens):
                atoms.append((piece, section.label, estimate_tokens(piece)))

    chunks: list[DocumentChunk] = []
    current: list[tuple[str, str, int]] = []
    current_tokens = 0

    def flush() -> None:
        nonlocal current, current_tokens
        if not current:
            return
        sources = tuple(dict.fromkeys(source for _, source, _ in current))
        chunks.append(
            DocumentChunk(
                index=len(chunks),
                text="\n\n".join(text for text, _, _ in current),
                sources=sources,
                token_count=current_tokens,
            )
        )
        overlap: list[tuple[str, str, int]] = []
        overlap_size = 0
        for atom in reversed(current):
            if atom[2] > overlap_tokens:
                break
            if overlap and overlap_size + atom[2] > overlap_tokens:
                break
            overlap.insert(0, atom)
            overlap_size += atom[2]
            if overlap_size >= overlap_tokens:
                break
        current = overlap
        current_tokens = overlap_size

    for atom in atoms:
        if current and current_tokens + atom[2] > target_tokens:
            flush()
        if current and current_tokens + atom[2] > target_tokens:
            current = []
            current_tokens = 0
        current.append(atom)
        current_tokens += atom[2]
    flush()
    return chunks


def _features(text: str) -> Counter[str]:
    normalized = unicodedata.normalize("NFKC", text).casefold()
    words = re.findall(r"[^\W_]+", normalized, flags=re.UNICODE)
    features: Counter[str] = Counter(f"word:{word}" for word in words)
    compact = "".join(char for char in normalized if char.isalnum())
    features.update(
        f"char:{compact[i:i + 3]}" for i in range(max(0, len(compact) - 2))
    )
    return features


def _cosine(left: Counter[str], right: Counter[str]) -> float:
    if not left or not right:
        return 0.0
    common = left.keys() & right.keys()
    numerator = sum(left[key] * right[key] for key in common)
    left_norm = math.sqrt(sum(value * value for value in left.values()))
    right_norm = math.sqrt(sum(value * value for value in right.values()))
    return numerator / (left_norm * right_norm) if left_norm and right_norm else 0.0


def _document_digest(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return hashlib.sha256(ext.encode() + b"\0" + data).hexdigest()


def get_document(filename: str, data: bytes) -> DocumentIndex:
    digest = _document_digest(filename, data)
    cached = _DOCUMENT_CACHE.get(digest)
    if cached is not None:
        _DOCUMENT_CACHE.move_to_end(digest)
        print(
            f"[timing] event=document_cache ms=0 hit=true "
            f"chunks={len(cached.chunks)} tokens={cached.token_count}"
        )
        return cached

    started = time.perf_counter()
    extraction_started = time.perf_counter()
    sections = [section for section in _read_sections(filename, data) if section.text]
    _log_timing("document_extraction", extraction_started, bytes=len(data), sections=len(sections))
    chunk_started = time.perf_counter()
    chunks = chunk_sections(sections)
    _log_timing(
        "document_chunking",
        chunk_started,
        chunks=len(chunks),
        tokens=sum(chunk.token_count for chunk in chunks),
    )
    vector_started = time.perf_counter()
    vectors = [_features(chunk.text) for chunk in chunks]
    _log_timing("document_vectors", vector_started, chunks=len(chunks))
    index = DocumentIndex(
        filename=filename,
        digest=digest,
        sections=sections,
        chunks=chunks,
        vectors=vectors,
    )
    _DOCUMENT_CACHE[digest] = index
    _DOCUMENT_CACHE.move_to_end(digest)
    while len(_DOCUMENT_CACHE) > MAX_CACHED_DOCUMENTS:
        _DOCUMENT_CACHE.popitem(last=False)
    _log_timing(
        "document_index",
        started,
        cache_hit="false",
        bytes=len(data),
        sections=len(sections),
        chunks=len(chunks),
        tokens=index.token_count,
    )
    return index


def clear_document_cache() -> None:
    _DOCUMENT_CACHE.clear()


def extract_text(filename: str, data: bytes) -> str:
    """Extract full text for compatibility; large AI reads use chunk tools."""
    sections = _read_sections(filename, data)
    if not sections:
        return f"[unsupported file type: {filename}]"
    text = "\n\n".join(section.text for section in sections if section.text).strip()
    return text or f"[no extractable text in {filename}]"


def retrieve_chunks(
    document: DocumentIndex,
    query: str,
    max_chunks: int = DEFAULT_RETRIEVAL_CHUNKS,
) -> list[DocumentChunk]:
    started = time.perf_counter()
    limit = min(max(int(max_chunks), 1), MAX_RETRIEVAL_CHUNKS)
    query_vector = _features(query)
    normalized_query = unicodedata.normalize("NFKC", query).casefold().strip()
    scored = []
    for index, (chunk, vector) in enumerate(zip(document.chunks, document.vectors)):
        score = _cosine(query_vector, vector)
        if normalized_query and normalized_query in chunk.text.casefold():
            score += 1.0
        scored.append((score, index))
    scored.sort(reverse=True)

    if not scored or scored[0][0] <= 0:
        _log_timing(
            "document_retrieval",
            started,
            chunks_total=len(document.chunks),
            chunks_selected=0,
            tokens=0,
        )
        return []

    selected: list[int] = []
    for score, index in scored:
        if score <= 0 and selected:
            break
        for candidate in (index, index - 1, index + 1):
            if 0 <= candidate < len(document.chunks) and candidate not in selected:
                selected.append(candidate)
                if len(selected) == limit:
                    break
        if len(selected) == limit:
            break
    selected.sort()
    result = [document.chunks[index] for index in selected]
    _log_timing(
        "document_retrieval",
        started,
        chunks_total=len(document.chunks),
        chunks_selected=len(result),
        tokens=sum(chunk.token_count for chunk in result),
    )
    return result


def _group_by_token_budget(items: list[str], budget: int) -> list[list[str]]:
    groups: list[list[str]] = []
    current: list[str] = []
    tokens = 0
    for item in items:
        item_tokens = estimate_tokens(item)
        if current and tokens + item_tokens > budget:
            groups.append(current)
            current = []
            tokens = 0
        current.append(item)
        tokens += item_tokens
    if current:
        groups.append(current)
    return groups


async def summarize_document(
    document: DocumentIndex,
    complete: Completion,
    detail: str = "standard",
) -> str:
    detail = detail if detail in {"brief", "standard", "detailed"} else "standard"
    cached = document.summaries.get(detail)
    if cached is not None:
        print(
            f"[timing] event=document_summary_cache ms=0 hit=true "
            f"chunks={len(document.chunks)} detail={detail}"
        )
        return cached
    if not document.chunks:
        return f"[no extractable text in {document.filename}]"

    started = time.perf_counter()
    map_output_tokens = {"brief": 140, "standard": 260, "detailed": 420}[detail]
    final_output_tokens = {"brief": 300, "standard": 650, "detailed": 1100}[detail]
    notes = []
    map_system = (
        "Create compact factual notes from one document chunk. Preserve all "
        "important names, dates, numbers, conditions, and warnings. Do not "
        "translate; keep the source language. Do not add facts. Include the "
        "provided source label with every point. Output notes only."
    )
    for chunk in document.chunks:
        note = await complete(
                map_system,
                f"Source: {chunk.citation}\n\n{chunk.text}",
                map_output_tokens,
            )
        notes.append(
            f"[source: {chunk.citation}]\n{note}"
        )

    level = 0
    reduce_system = (
        "Merge document notes into a faithful, non-redundant summary. Preserve "
        "names, dates, numbers, conditions, warnings, and source labels. Do not "
        "translate and do not add facts. Output the merged summary only."
    )
    while len(notes) > 1:
        groups = _group_by_token_budget(notes, REDUCE_INPUT_TOKENS)
        output_limit = final_output_tokens if len(groups) == 1 else map_output_tokens * 2
        notes = [
            await complete(
                reduce_system,
                "\n\n---\n\n".join(group),
                output_limit,
            )
            for group in groups
        ]
        level += 1

    all_sources = tuple(
        dict.fromkeys(source for chunk in document.chunks for source in chunk.sources)
    )
    summary = f"{notes[0]}\n\nSources covered: {', '.join(all_sources)}"
    document.summaries[detail] = summary
    _log_timing(
        "document_summary",
        started,
        cache_hit="false",
        chunks=len(document.chunks),
        reduce_levels=level,
        detail=detail,
    )
    return summary


def build_attachment_tools(
    attachments: dict[str, tuple[str, bytes]],
    complete: Completion | None = None,
) -> dict[str, Callable]:
    def resolve(attachment_id: str) -> DocumentIndex | None:
        attachment = attachments.get(str(attachment_id))
        return get_document(*attachment) if attachment is not None else None

    def read_attached_file(attachment_id: str) -> str:
        document = resolve(attachment_id)
        if document is None:
            return f"attachment is not available: {attachment_id}"
        if not document.chunks:
            return f"[no extractable text in {document.filename}]"
        if document.token_count > DIRECT_READ_MAX_TOKENS:
            return (
                "document is too large for a full read: "
                f"chunks={len(document.chunks)}, estimated_tokens={document.token_count}. "
                "Use summarize_attachment for a whole-document summary or "
                "search_attachment for a specific question."
            )
        text = "\n\n".join(section.text for section in document.sections if section.text)
        return f"[attached file: {document.filename}]\n{text}"

    def search_attachment(
        attachment_id: str,
        query: str,
        max_chunks: int = DEFAULT_RETRIEVAL_CHUNKS,
    ) -> str:
        document = resolve(attachment_id)
        if document is None:
            return f"attachment is not available: {attachment_id}"
        chunks = retrieve_chunks(document, query, max_chunks)
        if not chunks:
            return "no relevant text found in the attachment"
        return "\n\n".join(
            f"[source: {chunk.citation}; chunk: {chunk.index + 1}]\n{chunk.text}"
            for chunk in chunks
        )

    async def summarize_attachment(
        attachment_id: str,
        detail: str = "standard",
    ) -> str:
        document = resolve(attachment_id)
        if document is None:
            return f"attachment is not available: {attachment_id}"
        if complete is None:
            return "document summarization is unavailable in this request"
        return await summarize_document(document, complete, detail)

    return {
        "read_attached_file": read_attached_file,
        "search_attachment": search_attachment,
        "summarize_attachment": summarize_attachment,
    }


def build_read_attached_file(
    attachments: dict[str, tuple[str, bytes]],
) -> Callable[[str], str]:
    """Compatibility wrapper around the request-scoped attachment registry."""
    return build_attachment_tools(attachments)["read_attached_file"]
