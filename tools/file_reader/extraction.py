import io

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import SourceSection


SUPPORTED_EXTENSIONS = {
    "pdf",
    "docx",
    "xlsx",
    "txt",
    "md",
    "csv",
    "py",
    "js",
    "json",
}
TEXT_EXTENSIONS = {"txt", "md", "csv", "py", "js", "json"}


def read_sections(filename: str, data: bytes) -> list[SourceSection]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return [
            SourceSection(f"page {number}", (page.extract_text() or "").strip())
            for number, page in enumerate(reader.pages, start=1)
        ]
    if ext == "docx":
        document = Document(io.BytesIO(data))
        text = "\n\n".join(
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        return [SourceSection("document", text)]
    if ext == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections = []
        for sheet in workbook.worksheets:
            lines = [
                ", ".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sections.append(SourceSection(f"sheet {sheet.title}", "\n".join(lines)))
        return sections
    if ext in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
        pages = text.split("\f")
        return [
            SourceSection(
                f"page {number}" if len(pages) > 1 else "document",
                page.strip(),
            )
            for number, page in enumerate(pages, start=1)
        ]
    return []


def extract_text(filename: str, data: bytes) -> str:
    """Extract full text for compatibility; large AI reads use chunk tools."""
    sections = read_sections(filename, data)
    if not sections:
        return f"[unsupported file type: {filename}]"
    text = "\n\n".join(section.text for section in sections if section.text).strip()
    return text or f"[no extractable text in {filename}]"
