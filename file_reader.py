import io

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_CHARS = 80000

SUPPORTED_EXTENSIONS = {"pdf", "docx", "xlsx", "txt", "md", "csv", "py", "js", "json"}

TEXT_EXTENSIONS = {"txt", "md", "csv", "py", "js", "json"}


def _read_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            lines.append(", ".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        text = _read_pdf(data)
    elif ext == "docx":
        text = _read_docx(data)
    elif ext == "xlsx":
        text = _read_xlsx(data)
    elif ext in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
    else:
        return f"[unsupported file type: {filename}]"

    text = text.strip()
    if not text:
        return f"[no extractable text in {filename}]"
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n...[truncated]"
    return text
